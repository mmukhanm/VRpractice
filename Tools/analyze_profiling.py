"""Build an Excel report from real profiling CSV files, or an explicitly empty template.

Requires openpyxl. No fabricated samples, no extrapolated maximum N, no NaN -> 0.
"""
import argparse
import csv
import math
from pathlib import Path
import statistics
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Series, Reference
from openpyxl.styles import Font, PatternFill

COUNTS = [0, 100, 250, 500, 1000, 2000]
SERIES = ['Baseline', 'SRPBatcher', 'GPUInstancing', 'LOD']
METRICS = ['frame_ms', 'cpu_main_ms', 'render_thread_ms', 'gpu_ms', 'batches', 'setpass', 'tris', 'verts', 'gc_alloc_bytes', 'used_memory_bytes', 'physics_simulate_ms']

def finite(value):
    try:
        number = float(value)
        return number if math.isfinite(number) else None
    except (ValueError, TypeError):
        return None

def summarize(path):
    with Path(path).open(encoding='utf-8-sig', newline='') as stream:
        rows = list(csv.DictReader(stream))
    if not rows or 'frame_ms' not in rows[0]:
        return None
    result = {key: rows[0].get(key, '') for key in ['label', 'stage', 'checkpoint', 'N', 'series', 'behaviour', 'physics_case', 'lighting_case']}
    result['frames'] = len(rows)
    result['complete'] = len(rows) >= 300 and all(row.get('complete') == '1' for row in rows)
    for metric in METRICS:
        values = [v for row in rows if (v := finite(row.get(metric))) is not None]
        result[metric + '_mean'] = statistics.mean(values) if values else None
        result[metric + '_max'] = max(values) if values else None
        result[metric + '_valid_frames'] = len(values)
    # Physics captures may contain both phases. Keep them separate, never label a mixed mean "falling".
    for phase, condition in [('falling', lambda n: n > 0), ('sleeping', lambda n: n == 0)]:
        phase_rows = [r for r in rows if finite(r.get('awake_bodies')) is not None and condition(float(r['awake_bodies']))]
        result[phase + '_frames'] = len(phase_rows)
        for metric in ['frame_ms', 'physics_simulate_ms']:
            values = [v for row in phase_rows if (v := finite(row.get(metric))) is not None]
            result[phase + '_' + metric] = statistics.mean(values) if values else None
    result['path'] = str(path)
    return result

def graphics_results(summaries):
    result = {}
    for item in summaries:
        if item['stage'] != '3' or item['checkpoint'] != 'CP2' or item['behaviour'] != 'None' or not item['complete'] or item['frame_ms_valid_frames'] < 300:
            continue
        count = finite(item['N'])
        if count in COUNTS and item['series'] in SERIES:
            result[(item['series'], int(count))] = item
    return result

def maximum_counts(results, series, budget=13.9):
    items = [results.get((series, count)) for count in COUNTS]
    if any(item is None or item['frame_ms_valid_frames'] < 300 for item in items):
        return None, None, 'Недостаточно данных: нужны все 6 N и ≥300 кадров'
    mean_fit = [count for count, item in zip(COUNTS, items) if item['frame_ms_mean'] <= budget]
    max_fit = [count for count, item in zip(COUNTS, items) if item['frame_ms_max'] <= budget]
    return max(mean_fit, default='Нет N'), max(max_fit, default='Нет N'), 'Полная серия; только измеренные N, без интерполяции'

def workbook(summaries):
    book = Workbook()
    notes = book.active; notes.title = 'Инструкция'
    for row in [
        ['Андижан Мухаммед — профилирование Unity 6000.6.0f1'],
        ['Данные', 'Шаблон без замеров' if not summaries else 'Только фактические CSV; неполные записи не входят в график'],
        ['Бюджет, мс', 13.9],
        ['Среда', 'Укажите CPU/GPU, разрешение, API, устройство, частоту, настройки качества; см. environment.txt'],
        ['Frame vs CPU', 'frame_ms = unscaledDeltaTime; CPU Main/Render/GPU — отдельные счетчики Profiler'],
        ['N/A', 'Пустое поле = счетчик недоступен/не измерен; это не 0'],
        ['Повторные записи', 'Для одной пары серия/N берется последняя полная запись по времени файла'],
        ['Статистика', 'Среднее и максимум; выбросы не удаляются'],
        ['Deep Profile', 'Не включайте Deep Profile для основных серий и Bad/Good Compare; отдельную глубокую запись храните отдельно'],
        ['Правила', 'Не сравнивайте разное разрешение, положение/направление камеры, фоновые задачи и качество'],
    ]: notes.append(row)
    graphics = book.create_sheet('Графика')
    headers = ['N']
    for series in SERIES: headers += [series + '_' + name for name in ['mean_ms', 'max_ms', 'batches', 'setpass', 'tris', 'frames']]
    headers += ['Бюджет_ms']
    graphics.append(headers)
    results = graphics_results(summaries)
    for count in COUNTS:
        row = [count]
        for series in SERIES:
            item = results.get((series, count))
            row += [item.get(key) if item else None for key in ['frame_ms_mean', 'frame_ms_max', 'batches_mean', 'setpass_mean', 'tris_mean', 'frames']]
        row += [13.9]; graphics.append(row)
    chart = ScatterChart(); chart.scatterStyle = 'lineMarker'; chart.title = 'Время кадра от N'; chart.x_axis.title = 'N объектов'; chart.y_axis.title = 'Среднее время кадра, мс'
    chart.display_blanks = 'gap'; chart.width = 26; chart.height = 14
    for column in [2, 8, 14, 20, 26]:
        series = Series(Reference(graphics, min_col=column, min_row=1, max_row=7), Reference(graphics, min_col=1, min_row=2, max_row=7), title_from_data=True)
        if column == 26: series.graphicalProperties.line.solidFill = 'D63031'
        chart.series.append(series)
    graphics.add_chart(chart, 'A10')
    limits = book.create_sheet('Максимальное N')
    limits.append(['Серия', 'Max N по среднему ≤13.9', 'Max N по максимуму ≤13.9', 'Достаточность'])
    for series in SERIES: limits.append([series, *maximum_counts(results, series)])
    raw = book.create_sheet('Сводка CSV')
    fields = ['label','stage','checkpoint','N','series','behaviour','physics_case','lighting_case','complete','frames']
    fields += [metric + suffix for metric in METRICS for suffix in ['_mean','_max','_valid_frames']]
    fields += ['falling_frames','falling_frame_ms','falling_physics_simulate_ms','sleeping_frames','sleeping_frame_ms','sleeping_physics_simulate_ms','path']
    raw.append(fields)
    for item in summaries: raw.append([item.get(key) for key in fields])
    baseline = book.create_sheet('Этап 2')
    baseline.append(['Точка','CPU Main mean ms','CPU Main max ms','Render mean ms','Batches','SetPass','Tris','Verts','GC/frame B','Used Memory B','CPU/GPU вывод по Timeline'])
    for cp in ['CP1','CP2','CP3']:
        item = next((x for x in reversed(summaries) if x['stage']=='2' and x['checkpoint']==cp and x['complete']), {})
        baseline.append([cp, *[item.get(k) for k in ['cpu_main_ms_mean','cpu_main_ms_max','render_thread_ms_mean','batches_mean','setpass_mean','tris_mean','verts_mean','gc_alloc_bytes_mean','used_memory_bytes_mean']], None])
    hierarchy = book.create_sheet('CP3 Top5')
    hierarchy.append(['Ранг Total','Строка Hierarchy','Total ms','Self ms','Calls','GC Alloc B','Что означает','Скриншот'])
    for rank in range(1,6): hierarchy.append([rank])
    scripts = book.create_sheet('Этап 4 метод')
    scripts.append(['Метод','Total ms','Self ms','Calls','GC Alloc метода B/frame','GC.Collect интервал с','Compare / Screenshot'])
    scripts.append(['BadBehaviour.Update']); scripts.append(['GoodBehaviour.CheckDistance'])
    scripts.append(['Значения этого листа берутся из Hierarchy/Profile Analyzer, не из общих GC счетчиков CSV.'])
    physics = book.create_sheet('Этап 5 физика')
    physics.append(['Коллайдер','Фаза','Измерено кадров фазы','Frame mean ms','Physics.Simulate marker mean ms','Physics module ms (Profiler)','Вывод'])
    for collider in ['Convex','Box']:
        for phase in ['falling','sleeping']:
            valid = [x for x in summaries if x['stage']=='5' and x['physics_case'].startswith(collider+'_') and x.get(phase+'_frames',0)>0 and x['complete']]
            item = valid[-1] if valid else {}
            physics.append([collider,phase,item.get(phase+'_frames'),item.get(phase+'_frame_ms'),item.get(phase+'_physics_simulate_ms')])
    lighting = book.create_sheet('Этап 5 свет')
    lighting.append(['Случай','Frame mean ms','Render mean ms','GPU mean ms','Batches','Shadow passes (Frame Debugger)','Скриншот'])
    for case in ['Directional','RealtimeShadows','RealtimeNoShadows','Baked']:
        item = next((x for x in reversed(summaries) if x['stage']=='5' and x['lighting_case']==case and x['checkpoint']=='CP2' and x['complete']), {})
        lighting.append([case,*[item.get(k) for k in ['frame_ms_mean','render_thread_ms_mean','gpu_ms_mean','batches_mean']]])
    build = book.create_sheet('Этап 6')
    build.append(['Среда CP3','Frame mean ms','Frame max ms','CPU Main mean ms','Render mean ms','EditorLoop (Timeline)','Пояснение'])
    for stage,title in [('2','Editor'),('6','Development Build')]:
        item=next((x for x in reversed(summaries) if x['stage']==stage and x['checkpoint']=='CP3' and x['complete']),{})
        build.append([title,*[item.get(k) for k in ['frame_ms_mean','frame_ms_max','cpu_main_ms_mean','render_thread_ms_mean']]])
    for sheet in book:
        sheet.freeze_panes='B2'; sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]: cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='173F4F')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 24 if column[0].column>1 else 34
    return book

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input',type=Path,help='Capture root; keep separate Deep Profile runs outside this folder')
    parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args()
    paths=sorted(args.input.rglob('*.csv'),key=lambda p:(p.stat().st_mtime_ns,str(p))) if args.input else []
    if args.input and not paths: parser.error('No CSV files found; an empty directory is not a completed measurement.')
    summaries=[result for p in paths if (result:=summarize(p)) is not None]
    if args.input and not summaries: parser.error('No readable profiling captures found; no measurements to report.')
    args.output.parent.mkdir(parents=True,exist_ok=True)
    workbook(summaries).save(args.output)
    print(('EMPTY TEMPLATE: ' if not summaries else f'{len(summaries)} real captures: ')+str(args.output))

if __name__=='__main__': main()
